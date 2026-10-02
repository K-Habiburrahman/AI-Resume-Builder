import re
from pathlib import Path

from pypdf import PdfReader
from docx import Document


# ============================================================
# SECTION ALIASES
# ============================================================

SECTION_ALIASES = {
    "summary": {
        "summary",
        "professional summary",
        "profile",
        "career objective",
        "objective",
        "about me",
        "about",
    },

    "education": {
        "education",
        "academic background",
        "educational qualification",
        "educational qualifications",
        "academic qualifications",
        "academics",
    },

    "skills": {
        "skills",
        "technical skills",
        "technical skills & tools",
        "technical skills and tools",
        "skills & technologies",
        "skills and technologies",
        "technical competencies",
        "core skills",
        "technical expertise",
    },

    "projects": {
        "projects",
        "academic projects",
        "personal projects",
        "projects undertaken",
        "project experience",
        "key projects",
    },

    "experience": {
        "experience",
        "work experience",
        "professional experience",
        "work history",
        "employment history",
        "professional history",
        "career history",
        "employment",
        "internship",
        "internships",
    },

    "certifications": {
        "certifications",
        "certificates",
        "certification",
        "certifications & achievements",
        "certifications and achievements",
        "achievements",
        "licenses & certifications",
        "licenses and certifications",
    },
}


ALL_SECTION_NAMES = {
    alias
    for aliases in SECTION_ALIASES.values()
    for alias in aliases
}


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_whitespace(text):
    """
    Normalize spaces while preserving meaningful punctuation.
    """
    if not text:
        return ""

    text = text.replace("\xa0", " ")
    text = text.replace("\u200b", "")
    text = text.replace("\ufeff", "")

    # Collapse repeated whitespace
    text = re.sub(r"[ \t]+", " ", text)

    return text.strip()


def normalize_line(line):
    """
    Clean a single extracted line.
    """
    if not line:
        return ""

    line = line.replace("\xa0", " ")
    line = line.replace("\u200b", "")
    line = line.replace("\ufeff", "")

    # Remove markdown headings
    line = re.sub(r"^#+\s*", "", line)

    # Normalize bullet characters WITHOUT destroying hyphens inside text
    line = re.sub(r"^[•●▪◦○‣⁃]\s*", "", line)

    # Remove decorative bullets only when they are actually at the start
    line = re.sub(r"^\s*[-–—]\s+", "", line)

    # Normalize whitespace
    line = re.sub(r"\s+", " ", line)

    return line.strip()


# ============================================================
# PDF EXTRACTION
# ============================================================

def extract_text_from_pdf(file_path):
    """
    Extract text from PDF using pypdf.

    Important:
    PDF does not store text as normal paragraphs.
    Reading order may therefore be imperfect, especially
    for two-column resumes.
    """

    reader = PdfReader(file_path)

    pages = []

    for page in reader.pages:
        try:
            page_text = page.extract_text(
                extraction_mode="layout"
            )
        except TypeError:
            # Older pypdf versions
            page_text = page.extract_text()

        if page_text:
            pages.append(page_text)

    return "\n".join(pages)


# ============================================================
# DOCX EXTRACTION
# ============================================================

def extract_text_from_docx(file_path):
    """
    Extract paragraphs AND tables from DOCX.

    A lot of resumes use tables for:
    - contact information
    - education
    - skills
    - two-column layouts
    """

    document = Document(file_path)

    parts = []

    # Normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tables
    for table in document.tables:
        for row in table.rows:

            cells = []

            for cell in row.cells:
                cell_text = " ".join(
                    paragraph.text.strip()
                    for paragraph in cell.paragraphs
                    if paragraph.text.strip()
                )

                if cell_text:
                    cells.append(cell_text)

            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


# ============================================================
# TXT EXTRACTION
# ============================================================

def extract_text_from_txt(file_path):
    with open(
        file_path,
        "r",
        encoding="utf-8",
        errors="ignore"
    ) as file:
        return file.read()


# ============================================================
# UNIVERSAL EXTRACTION
# ============================================================

def extract_text(file_path):

    extension = Path(file_path).suffix.lower()

    if extension == ".pdf":
        return extract_text_from_pdf(file_path)

    if extension == ".docx":
        return extract_text_from_docx(file_path)

    if extension == ".txt":
        return extract_text_from_txt(file_path)

    raise ValueError(
        "Unsupported file format. Please upload PDF, DOCX or TXT."
    )


# ============================================================
# SECTION DETECTION
# ============================================================

def normalize_section_heading(text):
    """
    Convert:

        PROFESSIONAL SUMMARY:
        Skills & Technologies
        EDUCATION
        Projects -

    into a comparable heading.
    """

    text = normalize_line(text)

    text = text.lower()

    # Remove trailing punctuation
    text = re.sub(r"[\s:|•\-–—]+$", "", text)

    # Normalize ampersand
    text = text.replace("&", "and")

    # Normalize whitespace
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def find_section(line):

    normalized = normalize_section_heading(line)

    for section, aliases in SECTION_ALIASES.items():

        for alias in aliases:

            alias_normalized = normalize_section_heading(alias)

            if normalized == alias_normalized:
                return section

    return None


# ============================================================
# CONTACT INFORMATION
# ============================================================

def extract_email(text):

    match = re.search(
        r"\b[A-Za-z0-9._%+-]+"
        r"@[A-Za-z0-9.-]+\.[A-Za-z]{2,}\b",
        text
    )

    return match.group(0) if match else ""


def extract_phone(text):

    patterns = [

        # +91 9999999999
        r"(?:\+91[\s.-]?)?[6-9]\d{9}",

        # +91-99999-99999
        r"\+91[\s.-]?\d{5}[\s.-]?\d{5}",
    ]

    for pattern in patterns:

        match = re.search(pattern, text)

        if match:

            phone = re.sub(
                r"\D",
                "",
                match.group(0)
            )

            if phone.startswith("91") and len(phone) == 12:
                phone = phone[-10:]

            if len(phone) == 10:
                return phone

    return ""


def extract_contact_information(text):

    email = extract_email(text)
    phone = extract_phone(text)

    return email, phone


# ============================================================
# NAME EXTRACTION
# ============================================================

def looks_like_name(line):

    line = normalize_line(line)

    if not line:
        return False

    if "@" in line:
        return False

    if re.search(r"\d", line):
        return False

    if len(line) > 60:
        return False

    words = line.split()

    if len(words) < 2 or len(words) > 5:
        return False

    excluded = {
        "resume",
        "curriculum vitae",
        "cv",
        "profile",
        "objective",
        "summary",
        "education",
        "skills",
        "projects",
        "experience",
        "certifications",
    }

    if line.lower() in excluded:
        return False

    return True


def extract_name(text):

    lines = [
        normalize_line(line)
        for line in text.splitlines()
    ]

    lines = [
        line
        for line in lines
        if line
    ]

    # Usually the name occurs near the beginning
    for line in lines[:15]:

        if looks_like_name(line):
            return line

    return ""


# ============================================================
# LOCATION
# ============================================================

def extract_location(text):

    # Common Indian city/state patterns.
    # This intentionally remains conservative.
    locations = [
        "Mumbai",
        "Thane",
        "Pune",
        "Bengaluru",
        "Bangalore",
        "Hyderabad",
        "Delhi",
        "New Delhi",
        "Chennai",
        "Kolkata",
        "Ahmedabad",
        "Navi Mumbai",
    ]

    lower_text = text.lower()

    for location in locations:

        if location.lower() in lower_text:
            return location

    return ""


# ============================================================
# SECTION EXTRACTION
# ============================================================

def extract_sections(text):

    raw_lines = text.splitlines()

    lines = []

    for raw_line in raw_lines:

        cleaned = normalize_line(raw_line)

        if cleaned:
            lines.append(cleaned)

    sections = {
        "summary": [],
        "education": [],
        "skills": [],
        "projects": [],
        "experience": [],
        "certifications": [],
    }

    current_section = None

    for line in lines:

        detected_section = find_section(line)

        # ----------------------------------------------------
        # SECTION HEADER
        # ----------------------------------------------------

        if detected_section:

            current_section = detected_section
            continue

        # ----------------------------------------------------
        # CONTENT
        # ----------------------------------------------------

        if current_section:
            sections[current_section].append(line)

    # --------------------------------------------------------
    # Convert arrays to text
    # --------------------------------------------------------

    cleaned_sections = {}

    for section, content in sections.items():

        # Remove duplicate consecutive lines
        result = []

        previous = None

        for line in content:

            line = normalize_line(line)

            if not line:
                continue

            if line == previous:
                continue

            result.append(line)

            previous = line

        cleaned_sections[section] = "\n".join(result).strip()

    return cleaned_sections


# ============================================================
# RESUME PARSER
# ============================================================

def parse_resume(file_path):

    text = extract_text(file_path)

    if not text.strip():
        raise ValueError(
            "No readable text was found in the uploaded resume."
        )

    email, phone = extract_contact_information(text)

    name = extract_name(text)

    location = extract_location(text)

    sections = extract_sections(text)

    return {
        "name": name,
        "email": email,
        "phone": phone,
        "location": location,

        "summary": sections["summary"],
        "education": sections["education"],
        "skills": sections["skills"],
        "projects": sections["projects"],
        "experience": sections["experience"],
        "certifications": sections["certifications"],
    }